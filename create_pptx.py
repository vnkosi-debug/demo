from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Create presentation
prs = Presentation()

# Add title slide
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)

# Title
title = slide.shapes.title
title.text = "🚀 Agentic Protocol Engineering: The Wiring Harness of Trust"

# Subtitle
subtitle = slide.placeholders[1]
subtitle.text = "When Arrows Bleed... Trust Leaks Out"

# Add content boxes
left = Inches(0.5)
top = Inches(2)
width = Inches(4.5)
height = Inches(2)

# Problem section
txBox = slide.shapes.add_textbox(left, top, width, height)
tf = txBox.text_frame
tf.text = "🎯 The Production Reality Check"
p = tf.add_paragraph()
p.text = "Clean diagrams hide dangerous gaps in agentic systems"
p.level = 1

p = tf.add_paragraph()
p.text = "❌ What Fails in Production:"
p.level = 0

p = tf.add_paragraph()
p.text = "• ⏰ Timestamps vanish between retrieval and reasoning"
p.level = 1

p = tf.add_paragraph()
p.text = "• 🔍 Provenance stripped from critical data"
p.level = 1

p = tf.add_paragraph()
p.text = "• 🧠 Memory drifts without structured handoffs"
p.level = 1

p = tf.add_paragraph()
p.text = "• ⚡ Actions fire without governance checks"
p.level = 1

# Solution section
left = Inches(5.5)
txBox2 = slide.shapes.add_textbox(left, top, width, height)
tf2 = txBox2.text_frame
tf2.text = "✅ The Protocol Solution: Trust That Travels"

# Create table
rows = 5
cols = 3
table = slide.shapes.add_table(rows, cols, left, Inches(4.5), width, Inches(2)).table

# Set column widths
table.columns[0].width = Inches(1.2)
table.columns[1].width = Inches(1.8)
table.columns[2].width = Inches(1.5)

# Headers
table.cell(0, 0).text = "Protocol"
table.cell(0, 1).text = "What It Does"
table.cell(0, 2).text = "Why It Matters"

# Data
protocols = [
    ("🧩 MCP", "Model Context", "Stable, versioned context slots"),
    ("💬 ACP", "Agent Comm", "Traceable message envelopes"),
    ("🤝 A2A", "Agent-to-Agent", "Task coordination & state sync"),
    ("🌐 ANP", "Network", "Cross-org identity & policy")
]

for i, (proto, desc, benefit) in enumerate(protocols, 1):
    table.cell(i, 0).text = proto
    table.cell(i, 1).text = desc
    table.cell(i, 2).text = benefit

# Maturity ladder section
left = Inches(0.5)
top = Inches(6.5)
width = Inches(9)
height = Inches(1.5)

txBox3 = slide.shapes.add_textbox(left, top, width, height)
tf3 = txBox3.text_frame
tf3.text = "📈 Maturity Ladder: From Fragile to Bulletproof"
p = tf3.add_paragraph()
p.text = "🔴 L0: Ad Hoc Chaos → 🟠 L1: Sandbox Safety → 🟡 L2: Context Stability → 🟢 L3: Structured Comm → 🔵 L4: Federated Trust → 🟣 L5: Protocol Mesh"

# Demo section
top = Inches(8)
txBox4 = slide.shapes.add_textbox(left, top, width, Inches(1))
tf4 = txBox4.text_frame
tf4.text = "🎪 Live Demo: Protocol Envelope in Action"

# JSON demo
top = Inches(9)
txBox5 = slide.shapes.add_textbox(left, top, width, Inches(1.5))
tf5 = txBox5.text_frame
tf5.text = '''{
  "protocol": "ACP_Message",
  "metadata": {
    "id": "msg_1774875876",
    "session": "med_review_001",
    "timestamp": "2026-03-30T15:04:36Z"
  },
  "payload": {
    "tool": "generate_report",
    "execution_class": "restricted",
    "provenance": {
      "origin": "verified_medical_db",
      "trust_score": 0.95,
      "crypto_signature": "a9577c809c691ff9"
    },
    "governance": {
      "requires_approval": true,
      "approver_role": "Chief_Medical_Officer"
    }
  }
}'''

# Key takeaway
top = Inches(10.5)
txBox6 = slide.shapes.add_textbox(left, top, width, Inches(1))
tf6 = txBox6.text_frame
tf6.text = '💡 Key Takeaway: APIs connect components. Protocols connect trust.'
p = tf6.add_paragraph()
p.text = '"The diagram is just a map. The protocols are the wiring harness that keeps the whole thing intact when it moves."'
p = tf6.add_paragraph()
p.text = "— Austin & Neter, Chapter 8"

# Save the presentation
prs.save('Agentic_Protocol_Presentation.pptx')
print("PowerPoint created: Agentic_Protocol_Presentation.pptx")